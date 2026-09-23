"""
Fills in eBay's returned "recommendations" file (what you get back after
uploading a /ebay export-batch CSV to Seller Hub and eBay's AI processes it -
see ebay_csv.py's docstring for the full 3-step workflow) with data this bot
already captured during Queue Review approval, so nobody has to retype
price/condition/format by hand for every item before the final re-upload.

eBay's own recommendations file uses a DIFFERENT template internally than
the one this bot uploads (Template=eBay-listings-template_EBAY_US, not
eBay-taxonomy-mapping-template_US) - one sheet per matched category, named
"Cat-<CategoryName>...", each with its own #INFO/header preamble (matching
this bot's own upload format's "#INFO rows before the header" convention)
and a full classic-style column set (Start price, Quantity, Condition ID,
Format, Duration, Location, Shipping service 1 option, C:<Aspect> columns,
etc). eBay's own AI only ever seems to suggest category/title/description/
aspects - price/condition/format/location/shipping are left blank for a
human to fill in, which is exactly what this module automates from data
already on file: ebay_listing_data.price/condition_id/listing_format/
auction_duration (see EbayListingModal, item_flow.py), matched back to a
row via its "Custom Label (SKU)" cell (this bot's own pallet-N-item-M
format - see ebay_csv.custom_label).

Never overwrites a cell that already has something in it - whatever eBay's
own AI suggested, or anything a human already typed in, is left alone.
Location and Shipping service 1 option are only filled if EBAY_ITEM_LOCATION/
EBAY_SHIPPING_SERVICE are set in .env - both purely optional convenience
values used ONLY here, never required or blocking (unlike this bot's old
classic-template export, which used to hard-require them before every
export - see ebay_csv.py's docstring for why that no longer applies).

A real eBay-generated file has been observed with an invalid font
<family val="34"/> in its style/comment XML (valid range is 0-14) - this
makes openpyxl (and any strict OOXML reader) refuse to open the file at all
with "Max value is 14". _sanitize_workbook_zip() works around this by
clamping any out-of-range family value before handing the file to openpyxl -
purely cosmetic metadata, doesn't touch any actual data.

Known limitation: openpyxl drops "extended" (x14) data validation on save
(Excel dropdown pick-lists driven by the Categories sheet) - the filled
file still has all the same data, cells just may not have their dropdown
UI anymore. Macros (VBA) ARE preserved (keep_vba=True).
"""
import io
import os
import re
import tempfile
import warnings
import zipfile
from pathlib import Path

import openpyxl

import config
import database as db

SKU_PATTERN = re.compile(r"^pallet-(\d+)-item-(\d+)$")

# Column name -> whether it's filled from ebay_listing_data (per item) or
# from an optional config value (same for every item). Order doesn't matter -
# _fill_value_for_column decides per column.
FILLABLE_COLUMNS = (
    "Start price", "Quantity", "Condition ID", "Format", "Duration",
    "Location", "Shipping service 1 option",
)

_MAX_HEADER_SCAN_ROWS = 10


class RecommendationsFileError(Exception):
    """Raised when the uploaded file isn't a readable eBay recommendations workbook."""


class FillSummary:
    def __init__(self):
        self.filled_count = 0          # rows that got at least one cell filled
        self.sheet_count = 0           # Cat-* sheets that had at least one fill
        self.unmatched_skus = []       # SKUs present but not resolvable to a known item
        self.touched_location = False
        self.touched_shipping_service = False


def _sanitize_workbook_zip(raw_bytes: bytes) -> bytes:
    """
    Clamps any <family val="N"/> where N > 14 (outside OOXML's valid 0-14
    range) down to 2 (a generic sans-serif default) across every XML part -
    see this module's docstring for why eBay's own files can have this.
    """
    src = zipfile.ZipFile(io.BytesIO(raw_bytes))
    out_buffer = io.BytesIO()
    with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith(".xml"):
                text = data.decode("utf-8", errors="ignore")
                # Matches both spacing conventions: eBay's own generated
                # XML writes this with no space (<family val="34"/>), while
                # a workbook that's already passed through openpyxl once
                # writes one (<family val="2" />) - both need to match.
                fixed = re.sub(
                    r'<family val="(\d+)"\s*/>',
                    lambda m: '<family val="2"/>' if int(m.group(1)) > 14 else m.group(0),
                    text,
                )
                if fixed != text:
                    data = fixed.encode("utf-8")
            out.writestr(item, data)
    return out_buffer.getvalue()


def _find_header_row(ws) -> int | None:
    """The row containing a "Custom Label (SKU)" cell - eBay's own files put
    this on row 4 (after 3 #INFO rows) but this scans rather than assumes,
    in case a future export shifts it."""
    for row_idx in range(1, _MAX_HEADER_SCAN_ROWS + 1):
        for cell in ws[row_idx]:
            if cell.value == "Custom Label (SKU)":
                return row_idx
    return None


def _fill_value_for_column(column_name: str, listing: dict) -> str | None:
    if column_name == "Start price":
        return f"{listing['price']:.2f}" if listing.get("price") is not None else None
    if column_name == "Quantity":
        return "1"
    if column_name == "Condition ID":
        return listing.get("condition_id")
    if column_name == "Format":
        return listing.get("listing_format")
    if column_name == "Duration":
        return listing.get("auction_duration") if listing.get("listing_format") == "Auction" else None
    if column_name == "Location":
        return config.EBAY_ITEM_LOCATION or None
    if column_name == "Shipping service 1 option":
        return config.EBAY_SHIPPING_SERVICE or None
    return None


def fill_recommendations_file(raw_bytes: bytes) -> tuple:
    """
    Fills in every recognized item's price/quantity/condition/format/
    duration (and location/shipping service, if configured) across every
    "Cat-*" sheet in eBay's recommendations file, using data captured during
    Queue Review approval. Returns (Path to the filled .xlsm, FillSummary).
    Raises RecommendationsFileError if the file can't be opened at all.
    """
    try:
        sanitized = _sanitize_workbook_zip(raw_bytes)
        with warnings.catch_warnings():
            # openpyxl drops Excel's "extended" (x14) data validation on
            # save (dropdown pick-lists) - expected and disclosed to the
            # user in the command's response, not worth a console warning
            # on every use.
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
            wb = openpyxl.load_workbook(io.BytesIO(sanitized), keep_vba=True)
    except RecommendationsFileError:
        raise
    except Exception as e:
        raise RecommendationsFileError(f"couldn't open this as an Excel file ({e})") from e

    summary = FillSummary()

    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("Cat-"):
            continue
        ws = wb[sheet_name]
        header_row = _find_header_row(ws)
        if header_row is None:
            continue

        columns = {}
        for cell in ws[header_row]:
            if cell.value:
                columns[str(cell.value).strip()] = cell.column
        sku_col = columns.get("Custom Label (SKU)")
        if not sku_col:
            continue

        sheet_touched = False
        for row_idx in range(header_row + 1, ws.max_row + 1):
            sku_value = ws.cell(row=row_idx, column=sku_col).value
            sku = sku_value.strip() if isinstance(sku_value, str) else ""
            if not sku:
                continue

            match = SKU_PATTERN.match(sku)
            item = db.get_item_by_pallet_and_number(int(match.group(1)), int(match.group(2))) if match else None
            listing = db.get_ebay_listing_data(item["id"]) if item else None
            if not listing:
                summary.unmatched_skus.append(sku)
                continue

            row_touched = False
            for column_name in FILLABLE_COLUMNS:
                col_idx = columns.get(column_name)
                if not col_idx:
                    continue
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value not in (None, ""):
                    continue  # never overwrite an existing value
                value = _fill_value_for_column(column_name, listing)
                if not value:
                    continue
                cell.value = value
                row_touched = True
                sheet_touched = True
                if column_name == "Location":
                    summary.touched_location = True
                elif column_name == "Shipping service 1 option":
                    summary.touched_shipping_service = True
            if row_touched:
                summary.filled_count += 1

        if sheet_touched:
            summary.sheet_count += 1

    fd, path = tempfile.mkstemp(suffix=".xlsm", prefix="ebay_recommendations_filled_")
    os.close(fd)
    out_path = Path(path)
    wb.save(out_path)
    return out_path, summary
